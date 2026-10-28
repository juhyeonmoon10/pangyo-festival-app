"""Validate the official workbook locally. Never print student records or contact data."""
import argparse
import collections
import hashlib
import json
import re
import unicodedata
from pathlib import Path

import openpyxl


def text(value):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return unicodedata.normalize('NFKC', str(value or '')).strip()


def club_key(value):
    return re.sub(r'\s+', '', re.sub(r'\([^)]*\)', '', text(value)))


def analyze(workbook, catalog, aliases=None, approved_mappings=None):
    aliases = aliases or {}
    approved_mappings = approved_mappings or {}
    canonical = lambda value: aliases.get(club_key(value), club_key(value))
    final, grades = collections.defaultdict(list), collections.defaultdict(list)
    issues = []
    sheet = workbook.worksheets[0]
    rows = list(sheet.values)
    if len(rows) < 13 or text(rows[4][1]) != '학번' or text(rows[11][1]) != '학번':
        raise ValueError('UNRECOGNIZED_FINAL_SHEET')

    def add(target, sid, name, affiliation, cell):
        sid, name = text(sid), text(name)
        if not sid and not name:
            return
        if not re.fullmatch(r'[1-3][0-9]{4}', sid) or not name or len(name) > 60:
            issues.append({'cell': cell, 'code': 'INVALID_STUDENT_ROW'})
            return
        target[sid].append((name, canonical(affiliation), cell))

    for col in range(1, len(rows[4]), 2):
        if text(rows[4][col]) != '학번' or text(rows[4][col + 1]) != '이름':
            continue
        affiliation = rows[1][col]
        # Rows 6/7 are club leaders; rows 8-12 are capacity, room and headings.
        for index in [5, 6, *range(12, len(rows))]:
            add(final, rows[index][col], rows[index][col + 1], affiliation,
                f'{sheet.title}!{openpyxl.utils.get_column_letter(col + 1)}{index + 1}')
    for sheet in workbook.worksheets:
        if sheet.title not in ('1학년(반번호순)', '2학년(반번호순)'):
            continue
        for row, values in enumerate(sheet.iter_rows(min_row=2, max_col=3, values_only=True), 2):
            if any(values):
                add(grades, values[0], values[1], values[2], f'{sheet.title}!A{row}')

    by_club = collections.defaultdict(list)
    for key in catalog:
        by_club[canonical(key)].append(key)
    records, pending = [], []
    for sid in sorted(final.keys() | grades.keys()):
        a, b = final[sid], grades[sid]
        names = {r[0] for r in a + b}
        clubs = sorted({r[1] for r in a + b})
        reasons = []
        if len(names) != 1:
            reasons.append('NAME_CONFLICT')
        if not a or not b:
            reasons.append('SINGLE_SOURCE')
        elif {(r[0], r[1]) for r in a} != {(r[0], r[1]) for r in b}:
            reasons.append('SHEET_CONFLICT')
        keys = set()
        for club in clubs:
            mapped = approved_mappings.get(club, by_club.get(club))
            # An explicit [] is an approved non-participating club; absence is unknown.
            if mapped is None:
                reasons.append('UNMAPPED_CLUB')
            elif not isinstance(mapped, list) or any(key not in catalog for key in mapped):
                raise ValueError('INVALID_APPROVED_MAPPING')
            else:
                keys.update(mapped)
        cells = sorted({r[2] for r in a + b})
        reasons = sorted(set(reasons))
        records.append({'studentNumber': sid, 'realName': next(iter(names)) if len(names) == 1 else None,
                        'clubs': clubs, 'boothKeys': sorted(keys), 'status': 'review_required' if reasons else 'ready',
                        'issues': reasons, 'sourceCells': cells})
        if reasons:
            pending.append({'cells': cells, 'issues': reasons})
    summary = {'students': len(records), 'ready': sum(r['status'] == 'ready' for r in records),
               'pending': len(pending), 'clubs': len({c for r in records for c in r['clubs']}),
               'grades': dict(collections.Counter(r['studentNumber'][0] for r in records)),
               'reasonCounts': dict(collections.Counter(reason for r in records for reason in r['issues'])),
               'unmappedClubs': sorted({c for r in records for c in r['clubs'] if c not in approved_mappings and c not in by_club}),
               'exactDuplicatesCollapsed': sum(len(rows) - len({(r[0], r[1]) for r in rows}) for rows in final.values()),
               'invalidRows': issues, 'pendingCells': pending}
    return records, summary


def private_output(destination):
    root = Path(__file__).resolve().parents[1] / 'nfc-server' / 'private-output'
    resolved = destination.resolve()
    if not resolved.is_relative_to(root.resolve()) or resolved == root.resolve():
        raise ValueError('OUTPUT_MUST_BE_UNDER_IGNORED_PRIVATE_OUTPUT')
    return resolved


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('workbook', type=Path)
    parser.add_argument('--catalog', type=Path, required=True, help='JSON array of existing booth keys')
    parser.add_argument('--mapping', type=Path, help='Reviewed aliases and booth mappings, no student data')
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    output = private_output(args.output)
    if output.exists():
        raise ValueError('OUTPUT_ALREADY_EXISTS')
    catalog = json.loads(args.catalog.read_text(encoding='utf-8-sig'))
    if not isinstance(catalog, list) or not catalog or any(not isinstance(k, str) or not k for k in catalog) or len(set(catalog)) != len(catalog):
        raise ValueError('INVALID_CATALOG')
    mapping = json.loads(args.mapping.read_text(encoding='utf-8-sig')) if args.mapping else {}
    workbook = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True)
    try:
        records, summary = analyze(workbook, catalog, mapping.get('aliases'), mapping.get('clubs'))
    finally:
        workbook.close()
    if summary['invalidRows']:
        raise ValueError('MALFORMED_ROWS_REQUIRE_REVIEW')
    source_hash = hashlib.sha256(args.workbook.read_bytes()).hexdigest()
    payload = {'format': 1, 'sourceSha256': source_hash, 'catalogKeys': sorted(catalog), 'records': records}
    output.mkdir(parents=True)
    (output / 'roster.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
    (output / 'report.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k: v for k, v in summary.items() if k not in ('pendingCells', 'invalidRows')}, ensure_ascii=False))


if __name__ == '__main__':
    try:
        main()
    except (ValueError, OSError) as error:
        # File paths and records must not be echoed by CI or a public build.
        raise SystemExit(str(error) if isinstance(error, ValueError) else 'LOCAL_FILE_ERROR') from None

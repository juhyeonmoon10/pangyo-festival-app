import importlib.util
import unittest
from pathlib import Path
from openpyxl import Workbook

spec = importlib.util.spec_from_file_location('roster', Path(__file__).resolve().parents[1] / 'tools/import-review-roster.py')
roster = importlib.util.module_from_spec(spec)
spec.loader.exec_module(roster)


def fixture():
    w = Workbook()
    s = w.active
    s.title = '동아리학생명단(1,2학년)최종명단'
    for col, club in [(2, '테스트 동아리 (코딩)'), (4, '다른동아리')]:
        s.cell(2, col, club)
        for row in [5, 12]:
            s.cell(row, col, '학번')
            s.cell(row, col + 1, '이름')
    s.cell(6, 2, 11001)
    s.cell(6, 3, '가상학생')
    s.cell(13, 2, '11001')
    s.cell(13, 3, '가상학생')
    s.cell(4, 2, '가상교사: 수집 금지')
    s.cell(8, 2, '25명')
    g = w.create_sheet('1학년(반번호순)')
    g.append(['학번', '이름', '동아리'])
    g.append([11001, '가상학생', '테스트동아리'])
    return w


class ImportTests(unittest.TestCase):
    def test_deduplicates_leaders_and_ignores_teachers(self):
        records, report = roster.analyze(fixture(), ['테스트 동아리'])
        self.assertEqual(len(records), 1)
        self.assertEqual(report['exactDuplicatesCollapsed'], 1)
        self.assertEqual(records[0]['status'], 'ready')
        self.assertEqual(records[0]['boothKeys'], ['테스트 동아리'])
        self.assertNotIn('가상교사', str(records))

    def test_different_name_or_club_quarantined(self):
        w = fixture()
        w.worksheets[1].cell(2, 2, '다른학생')
        w.worksheets[1].cell(2, 3, '다른동아리')
        records, _ = roster.analyze(w, ['테스트 동아리', '다른동아리'])
        self.assertIsNone(records[0]['realName'])
        self.assertEqual(records[0]['status'], 'review_required')
        self.assertIn('SHEET_CONFLICT', records[0]['issues'])

    def test_unknown_club_not_treated_as_no_club(self):
        records, _ = roster.analyze(fixture(), ['다른동아리'])
        self.assertIn('UNMAPPED_CLUB', records[0]['issues'])

    def test_only_explicit_mapping_can_confirm_nonparticipation(self):
        records, _ = roster.analyze(fixture(), ['다른동아리'], approved_mappings={'테스트동아리': []})
        self.assertEqual(records[0]['status'], 'ready')

    def test_grade_only_and_single_source_pending(self):
        w = fixture()
        w.worksheets[1].append([11002, '가상둘', '다른동아리'])
        records, _ = roster.analyze(w, ['테스트 동아리', '다른동아리'])
        self.assertIn('SINGLE_SOURCE', records[1]['issues'])

    def test_same_name_different_numbers_not_merged(self):
        w = fixture()
        w.worksheets[1].append([11002, '가상학생', '다른동아리'])
        records, _ = roster.analyze(w, ['테스트 동아리'])
        self.assertEqual(len(records), 2)

    def test_invalid_student_number_detected(self):
        w = fixture()
        w.worksheets[0].cell(14, 2, 'bad')
        w.worksheets[0].cell(14, 3, '가상학생')
        _, report = roster.analyze(w, ['테스트 동아리'])
        self.assertEqual(len(report['invalidRows']), 1)

    def test_export_outside_private_output_rejected(self):
        with self.assertRaisesRegex(ValueError, 'IGNORED_PRIVATE_OUTPUT'):
            roster.private_output(Path('public/roster.json'))


if __name__ == '__main__':
    unittest.main()
